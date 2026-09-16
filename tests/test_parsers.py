"""
Tests for multi-format document parsers (PDF, DOCX, PPTX, TXT) in Prometheus.
"""
from __future__ import annotations

from pathlib import Path

import docx
import pymupdf
from pptx import Presentation

from prometheus.retrieval.parsers import (
    parse_document,
    parse_docx,
    parse_pdf,
    parse_pptx,
    parse_txt,
)


def test_parse_pdf_preserves_page_numbers(tmp_path: Path):
    # Create a 2-page PDF
    pdf_path = tmp_path / "test_doc.pdf"
    doc = pymupdf.open()
    
    page1 = doc.new_page()
    page1.insert_text((50, 50), "First page intro text.\n\nFirst page detailed paragraph.")
    
    page2 = doc.new_page()
    page2.insert_text((50, 50), "Second page conclusion text.")
    
    doc.save(str(pdf_path))
    doc.close()

    parsed = parse_pdf(pdf_path, document_id="pdf-doc-1")
    assert parsed.document_id == "pdf-doc-1"
    assert parsed.filename == "test_doc.pdf"
    assert parsed.file_type == "pdf"
    assert len(parsed.sections) >= 2

    # Check that page 1 and page 2 are tracked
    pages = [s.page_number for s in parsed.sections]
    assert 1 in pages
    assert 2 in pages
    assert any("First page" in s.text for s in parsed.sections if s.page_number == 1)
    assert any("Second page" in s.text for s in parsed.sections if s.page_number == 2)


def test_parse_docx_extracts_headings_and_tables(tmp_path: Path):
    docx_path = tmp_path / "test_doc.docx"
    doc = docx.Document()
    doc.add_heading("Project Overview", level=1)
    doc.add_paragraph("This is the main body paragraph describing the project.")
    
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Accuracy"
    table.cell(1, 1).text = "95%"
    
    doc.save(str(docx_path))

    parsed = parse_docx(docx_path, document_id="docx-doc-1")
    assert parsed.document_id == "docx-doc-1"
    assert parsed.file_type == "docx"

    section_types = [s.section_type for s in parsed.sections]
    assert "heading" in section_types
    assert "table" in section_types
    assert any("Project Overview" in s.text for s in parsed.sections)
    assert any("Accuracy | 95%" in s.text for s in parsed.sections)


def test_parse_pptx_preserves_slide_numbers(tmp_path: Path):
    pptx_path = tmp_path / "test_presentation.pptx"
    prs = Presentation()
    
    # Slide 1: Title
    slide_layout = prs.slide_layouts[0]
    slide1 = prs.slides.add_slide(slide_layout)
    slide1.shapes.title.text = "Introduction Slide"
    slide1.placeholders[1].text = "Subtitle content on slide 1"

    # Slide 2: Content
    slide_layout_content = prs.slide_layouts[1]
    slide2 = prs.slides.add_slide(slide_layout_content)
    slide2.shapes.title.text = "Architecture Details"
    slide2.placeholders[1].text = "Key architectural takeaways on slide 2"

    prs.save(str(pptx_path))

    parsed = parse_pptx(pptx_path, document_id="pptx-doc-1")
    assert parsed.document_id == "pptx-doc-1"
    assert parsed.file_type == "pptx"

    slides = [s.slide_number for s in parsed.sections]
    assert 1 in slides
    assert 2 in slides
    assert any("Introduction Slide" in s.text for s in parsed.sections if s.slide_number == 1)
    assert any("Architecture Details" in s.text for s in parsed.sections if s.slide_number == 2)


def test_parse_txt_handles_paragraphs_and_headings(tmp_path: Path):
    txt_path = tmp_path / "notes.txt"
    txt_path.write_text(
        "INTRODUCTION\n\nThis is a paragraph under the uppercase heading.\n\n1.1 Methodology\n\nMethod details go here.",
        encoding="utf-8",
    )

    parsed = parse_txt(txt_path, document_id="txt-doc-1")
    assert parsed.document_id == "txt-doc-1"
    assert parsed.file_type == "txt"
    assert len(parsed.sections) >= 2


def test_parse_document_dispatcher(tmp_path: Path):
    txt_path = tmp_path / "sample.txt"
    txt_path.write_text("Hello Prometheus RAG", encoding="utf-8")
    parsed = parse_document(txt_path)
    assert parsed.file_type == "txt"
    assert "Hello Prometheus RAG" in parsed.full_text
