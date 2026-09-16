"""
Multi-format document parsing module for Prometheus.

Supports:
- PDF (.pdf) via PyMuPDF, preserving 1-based page numbers.
- DOCX (.docx) via python-docx, preserving paragraph hierarchy and tables.
- PPTX (.pptx) via python-pptx, preserving 1-based slide numbers and tables.
- Plain text (.txt) with robust encoding fallback (UTF-8, latin-1).

All parsers return a unified ParsedDocument containing DocumentSection elements
with accurate location metadata (page_number or slide_number) so that downstream
chunking and citations have authentic references.
"""
from __future__ import annotations

import re
import uuid
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any


@dataclass
class DocumentSection:
    """A granular structural unit of extracted document text with location metadata."""

    text: str
    page_number: int | None = None
    slide_number: int | None = None
    heading: str | None = None
    section_type: str = "text"  # "text", "heading", "table", "list"
    metadata: dict[str, Any] = field(default_factory=dict)


@dataclass
class ParsedDocument:
    """Canonical representation of an ingested document regardless of original format."""

    document_id: str
    filename: str
    file_type: str
    sections: list[DocumentSection]
    title: str | None = None

    @property
    def full_text(self) -> str:
        return "\n\n".join(s.text for s in self.sections if s.text.strip())


SUPPORTED_EXTENSIONS = frozenset({".pdf", ".docx", ".pptx", ".txt"})


def parse_pdf(file_path: str | Path, document_id: str | None = None) -> ParsedDocument:
    """Extract text from a PDF file per-page, preserving 1-based page numbers."""
    import pymupdf

    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"PDF file not found: {path}")

    doc_id = document_id or str(uuid.uuid4())
    sections: list[DocumentSection] = []

    with pymupdf.open(str(path)) as doc:
        title = doc.metadata.get("title") if doc.metadata else None
        for page_idx, page in enumerate(doc, start=1):
            blocks = page.get_text("blocks", sort=True)
            for block in blocks:
                # PyMuPDF block tuple: (x0, y0, x1, y1, text, block_no, block_type)
                # block_type == 0 is text, block_type == 1 is image
                if len(block) > 6 and block[6] != 0:
                    continue
                block_text = block[4].strip() if len(block) > 4 and isinstance(block[4], str) else ""
                if not block_text:
                    continue

                # Preserve paragraph boundaries if a single block contains double-newlines
                paragraphs = [p.strip() for p in block_text.split("\n\n") if p.strip()]
                for p in paragraphs:
                    sections.append(
                        DocumentSection(
                            text=p,
                            page_number=page_idx,
                            section_type="text",
                        )
                    )

    if not sections:
        # Document might be empty or scanned images without OCR
        sections.append(
            DocumentSection(
                text="",
                page_number=1,
                section_type="text",
            )
        )

    return ParsedDocument(
        document_id=doc_id,
        filename=path.name,
        file_type="pdf",
        sections=sections,
        title=title or path.stem,
    )


def _extract_docx_heading_level(style_name: str) -> int | None:
    if not style_name:
        return None
    m = re.match(r"heading\s*(\d)", style_name.strip(), re.IGNORECASE)
    if m:
        return max(1, min(6, int(m.group(1))))
    if style_name.lower() in ("title", "subtitle"):
        return 1
    return None


def parse_docx(file_path: str | Path, document_id: str | None = None) -> ParsedDocument:
    """Extract text from a DOCX file, preserving headings, lists, and tables."""
    import docx

    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"DOCX file not found: {path}")

    doc_id = document_id or str(uuid.uuid4())
    word_doc = docx.Document(str(path))
    sections: list[DocumentSection] = []

    current_heading: str | None = None
    title = None
    if word_doc.core_properties and word_doc.core_properties.title:
        title = word_doc.core_properties.title.strip() or None

    body = word_doc.element.body
    for child in body:
        tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag

        if tag == "p":
            from docx.text.paragraph import Paragraph

            para = Paragraph(child, word_doc)
            raw_text = para.text.strip()
            if not raw_text:
                continue

            style_name = para.style.name if para.style else ""
            heading_lvl = _extract_docx_heading_level(style_name)

            if heading_lvl is not None:
                current_heading = raw_text
                sections.append(
                    DocumentSection(
                        text=raw_text,
                        heading=current_heading,
                        section_type="heading",
                        metadata={"level": heading_lvl},
                    )
                )
            else:
                sections.append(
                    DocumentSection(
                        text=raw_text,
                        heading=current_heading,
                        section_type="text",
                    )
                )

        elif tag == "tbl":
            from docx.table import Table

            tbl = Table(child, word_doc)
            rows = []
            for row in tbl.rows:
                cells = [c.text.strip() for c in row.cells]
                rows.append(" | ".join(cells))
            tbl_text = "\n".join(rows).strip()
            if tbl_text:
                sections.append(
                    DocumentSection(
                        text=tbl_text,
                        heading=current_heading,
                        section_type="table",
                    )
                )

    return ParsedDocument(
        document_id=doc_id,
        filename=path.name,
        file_type="docx",
        sections=sections,
        title=title or path.stem,
    )


def parse_pptx(file_path: str | Path, document_id: str | None = None) -> ParsedDocument:
    """Extract text from a PowerPoint PPTX file, preserving 1-based slide numbers."""
    from pptx import Presentation
    from pptx.enum.shapes import PP_PLACEHOLDER

    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"PPTX file not found: {path}")

    doc_id = document_id or str(uuid.uuid4())
    prs = Presentation(str(path))
    sections: list[DocumentSection] = []

    title = None
    if prs.core_properties and prs.core_properties.title:
        title = prs.core_properties.title.strip() or None

    for slide_idx, slide in enumerate(prs.slides, start=1):
        slide_title = None
        for shape in slide.shapes:
            if shape.is_placeholder:
                try:
                    is_title_type = shape.placeholder_format.type in (
                        PP_PLACEHOLDER.TITLE,
                        PP_PLACEHOLDER.CENTER_TITLE,
                        PP_PLACEHOLDER.VERTICAL_TITLE,
                    )
                    if is_title_type and shape.has_text_frame and shape.text_frame.text.strip():
                        slide_title = shape.text_frame.text.strip()
                        break
                except (AttributeError, KeyError, ValueError):
                    continue

        # If title found, add as heading section
        if slide_title:
            sections.append(
                DocumentSection(
                    text=slide_title,
                    slide_number=slide_idx,
                    heading=slide_title,
                    section_type="heading",
                )
            )

        # Extract other shapes
        for shape in slide.shapes:
            if shape.has_text_frame:
                frame_text = shape.text_frame.text.strip()
                if not frame_text or frame_text == slide_title:
                    continue
                sections.append(
                    DocumentSection(
                        text=frame_text,
                        slide_number=slide_idx,
                        heading=slide_title,
                        section_type="text",
                    )
                )
            elif shape.has_table:
                rows = []
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    rows.append(" | ".join(cells))
                tbl_text = "\n".join(rows).strip()
                if tbl_text:
                    sections.append(
                        DocumentSection(
                            text=tbl_text,
                            slide_number=slide_idx,
                            heading=slide_title,
                            section_type="table",
                        )
                    )

    return ParsedDocument(
        document_id=doc_id,
        filename=path.name,
        file_type="pptx",
        sections=sections,
        title=title or path.stem,
    )


def parse_txt(file_path: str | Path, document_id: str | None = None) -> ParsedDocument:
    """Extract text from a plain-text file with robust encoding fallbacks."""
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"TXT file not found: {path}")

    doc_id = document_id or str(uuid.uuid4())
    try:
        content = path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        try:
            content = path.read_text(encoding="latin-1")
        except Exception as exc:
            raise ValueError(f"Could not decode text file {path.name}: {exc}") from exc

    paragraphs = [p.strip() for p in re.split(r"\n{2,}", content) if p.strip()]
    sections: list[DocumentSection] = []

    current_heading: str | None = None
    for p in paragraphs:
        lines = p.splitlines()
        first_line = lines[0].strip() if lines else p.strip()
        # Simple heading heuristic: short line, all caps or numbered
        is_short = len(lines) == 1 and len(first_line.split()) <= 10
        is_heading = is_short and (
            first_line.isupper()
            or re.match(r"^\d+(\.\d+)*\s+[A-Z]", first_line)
            or re.match(r"^(chapter|section|part)\s+\d+", first_line, re.IGNORECASE)
        )

        if is_heading:
            current_heading = first_line
            sections.append(
                DocumentSection(
                    text=first_line,
                    heading=current_heading,
                    section_type="heading",
                )
            )
        else:
            sections.append(
                DocumentSection(
                    text=p,
                    heading=current_heading,
                    section_type="text",
                )
            )

    return ParsedDocument(
        document_id=doc_id,
        filename=path.name,
        file_type="txt",
        sections=sections,
        title=path.stem,
    )


def parse_document(file_path: str | Path, document_id: str | None = None) -> ParsedDocument:
    """Dispatch to appropriate parser based on file extension."""
    path = Path(file_path)
    ext = path.suffix.lower()

    if ext == ".pdf":
        return parse_pdf(path, document_id=document_id)
    elif ext == ".docx":
        return parse_docx(path, document_id=document_id)
    elif ext == ".pptx":
        return parse_pptx(path, document_id=document_id)
    elif ext == ".txt":
        return parse_txt(path, document_id=document_id)
    else:
        raise ValueError(
            f"Unsupported document format '{ext}'. Supported formats: {sorted(SUPPORTED_EXTENSIONS)}"
        )
